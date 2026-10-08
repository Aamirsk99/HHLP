"""Hindivine Food List: every food we have nutrition data for, in one table.

Sources
  1. IFCT 2017 (Indian Food Composition Tables, National Institute of Nutrition, ICMR) - 542 raw Indian foods, per 100 g
     (data/ifct2017.csv, from the MIT-licensed @ifct2017/compositions package).
  2. Hindivine ingredient catalogue (catalog.py) - cooking ingredients, per 100 g, with household measures.
  3. Hindivine recipe series - 5,165 dishes, per serving, with recipe numbers.

    python3 foodlist.py OUT_DIR      writes Hindivine-Food-List.xlsx, Hindivine-Food-List.csv and Hindivine-Food-Finder.html
"""
import csv
import json
import os
import sys

from catalog import CAT
from core import diet

HERE = os.path.dirname(os.path.abspath(__file__))
COLS = ['ID', 'Food', 'Hindi / local name', 'Category', 'Source', 'Diet', 'Basis', 'Household measure',
        'Energy (kcal)', 'Protein (g)', 'Carbs (g)', 'Fat (g)', 'Fibre (g)', 'Sugar (g)', 'Sat. fat (g)', 'Cholesterol (mg)',
        'Calcium (mg)', 'Iron (mg)', 'Sodium (mg)', 'Potassium (mg)', 'Vitamin C (mg)', 'Recipe no.', 'Best for']
NONVEG_GROUPS = {'Marine Fish', 'Animal Meat', 'Poultry', 'Fresh Water Fish and Shellfish', 'Marine Shellfish', 'Marine Mollusks'}
DIET_NAME = {'vegan': 'Vegan', 'veg': 'Vegetarian', 'egg': 'Egg', 'nonveg': 'Non-veg'}
CAT_NAME = {'spice': 'Spices & Condiments', 'veg': 'Vegetables', 'grain': 'Cereals, Millets & Flours', 'pulse': 'Pulses & Legumes',
            'fruit': 'Fruits', 'nut': 'Nuts & Seeds', 'cond': 'Sauces & Condiments', 'aro': 'Herbs & Aromatics', 'leafy': 'Leafy Vegetables',
            'dairy': 'Milk & Dairy', 'fat': 'Oils & Fats', 'meat': 'Meat & Poultry', 'fish': 'Fish & Seafood', 'sweet': 'Sugars & Sweeteners',
            'drink': 'Beverages', 'egg': 'Eggs'}


def num(x, scale=1.0, nd=1):
    try:
        v = float(x) * scale
    except (TypeError, ValueError):
        return None
    return round(v, nd)


def hindi(lang):
    for part in lang.split(';'):
        part = part.strip()
        if part.startswith('H.') or ', H.' in part[:8]:
            return part.split('.', 1)[1].strip().rstrip('.')
    return ''


def ifct_rows():
    out = []
    with open(os.path.join(HERE, 'data', 'ifct2017.csv'), encoding='utf-8') as fh:
        for r in csv.DictReader(fh):
            grp = r['grup']
            d = 'Non-veg' if grp in NONVEG_GROUPS else 'Egg' if grp == 'Egg and Egg Products' else 'Vegetarian'
            out.append({
                'ID': 'IF-' + r['code'], 'Food': r['name'], 'Hindi / local name': hindi(r['lang']), 'Category': grp,
                'Source': 'IFCT 2017 (NIN, ICMR)', 'Diet': d, 'Basis': 'per 100 g', 'Household measure': '',
                'Energy (kcal)': num(r['enerc'], 1 / 4.184, 0), 'Protein (g)': num(r['protcnt']), 'Carbs (g)': num(r['choavldf']),
                'Fat (g)': num(r['fatce']), 'Fibre (g)': num(r['fibtg']), 'Sugar (g)': num(r['fsugar']), 'Sat. fat (g)': num(r['fasat']),
                'Cholesterol (mg)': num(r['cholc'], 1000, 0), 'Calcium (mg)': num(r['ca'], 1000, 0), 'Iron (mg)': num(r['fe'], 1000, 1),
                'Sodium (mg)': num(r['na'], 1000, 0), 'Potassium (mg)': num(r['k'], 1000, 0), 'Vitamin C (mg)': num(r['vitc'], 1000, 0),
                'Recipe no.': '', 'Best for': '', '_local': r['lang']})
    return out


def ingredient_rows():
    out = []
    for i, (k, c) in enumerate(sorted(CAT.items(), key=lambda kv: kv[1]['name'].lower()), 1):
        if k in ('water', 'ice'):
            continue
        units = [u for u in c['units'] if u[0] != 'g']
        measure = '; '.join(f'1 {u} = {g:g} {"ml" if c["cat"] == "drink" or k in ("milk", "water", "coconutwater") else "g"}' for u, g in units)
        tags = c['tags']
        d = 'Non-veg' if tags & {'chicken', 'mutton', 'fish', 'shellfish'} else 'Egg' if 'egg' in tags else 'Vegetarian'
        out.append({'ID': f'IN-{i:03d}', 'Food': c['name'], 'Hindi / local name': '', 'Category': CAT_NAME.get(c['cat'], c['cat']),
                    'Source': 'Hindivine ingredient list', 'Diet': d, 'Basis': 'per 100 g', 'Household measure': measure,
                    'Energy (kcal)': round(c['kcal']), 'Protein (g)': c['p'], 'Carbs (g)': c['c'], 'Fat (g)': c['f'], 'Fibre (g)': c['fib'],
                    'Sugar (g)': None, 'Sat. fat (g)': None, 'Cholesterol (mg)': None, 'Calcium (mg)': None, 'Iron (mg)': None,
                    'Sodium (mg)': None, 'Potassium (mg)': None, 'Vitamin C (mg)': None, 'Recipe no.': '', 'Best for': ''})
    return out


def dish_rows():
    import build as B
    out = []
    for r in B.ordered(True):
        n = r['nut']
        out.append({'ID': f'HR-{r["no"]:04d}', 'Food': r['name'], 'Hindi / local name': '', 'Category': r['cat'] + (f' · {r["sub"]}' if r['sub'] else ''),
                    'Source': 'Hindivine recipe series', 'Diet': DIET_NAME[diet(r)], 'Basis': 'per serving',
                    'Household measure': f'1 serving (recipe serves {r["serves"]})',
                    'Energy (kcal)': round(n['kcal']), 'Protein (g)': round(n['p'], 1), 'Carbs (g)': round(n['c'], 1), 'Fat (g)': round(n['f'], 1),
                    'Fibre (g)': round(n['fib'], 1), 'Sugar (g)': None, 'Sat. fat (g)': None, 'Cholesterol (mg)': None, 'Calcium (mg)': None,
                    'Iron (mg)': None, 'Sodium (mg)': None, 'Potassium (mg)': None, 'Vitamin C (mg)': None,
                    'Recipe no.': r['no'], 'Best for': ', '.join(r['meals'])})
    return out


def all_rows():
    return ifct_rows() + ingredient_rows() + dish_rows()


# ---------------- outputs ----------------
def write_csv(rows, path):
    with open(path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=COLS, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    NAVY, SKY, CREAM = '0B3B66', '29A8E0', 'F2F7FB'
    head_font, head_fill = Font(name='Arial', bold=True, color='FFFFFF', size=10), PatternFill('solid', fgColor=NAVY)
    base, bold = Font(name='Arial', size=10), Font(name='Arial', size=10, bold=True)
    thin = Border(bottom=Side(style='thin', color='D6E7F3'))
    widths = {'ID': 10, 'Food': 42, 'Hindi / local name': 22, 'Category': 30, 'Source': 24, 'Diet': 11, 'Basis': 12,
              'Household measure': 30, 'Recipe no.': 10, 'Best for': 34}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Read Me'
    counts = {}
    for r in rows:
        counts[r['Source']] = counts.get(r['Source'], 0) + 1
    lines = [('Hindivine Food List', 'title'), ('www.hindivine.com · app@hindivine.com', 'sub'), ('', ''),
             (f'{len(rows):,} foods with energy, protein, carbohydrate, fat and fibre.', ''), ('', ''), ('Sources', 'h')]
    lines += [(f'{k}: {v:,} items', '') for k, v in counts.items()]
    lines += [('', ''), ('Sheets', 'h'),
              ('All Foods: every item, filterable by category, source and diet.', ''),
              ('Indian Foods (IFCT): raw Indian foods per 100 g, with sugar, saturated fat, cholesterol, calcium, iron, sodium, potassium and vitamin C.', ''),
              ('Hindivine Dishes: the 5,165 recipes, per serving; the recipe number matches the Hindivine Healthy Recipe Series.', ''),
              ('Ingredients: cooking ingredients per 100 g with household measures (cup, tbsp, tsp, pieces).', ''),
              ('Portion Calculator: type a food ID and an amount to get the nutrition for that portion.', ''),
              ('', ''), ('Notes', 'h'),
              ('Values for IFCT foods are for the raw, edible part per 100 g; energy was converted from kJ (÷ 4.184).', ''),
              ('Dish values are calculated from the recipe ingredients (IFCT 2017 / USDA) for one serving; they vary with brands and the oil used.', ''),
              ('Blank cells mean the value is not available in that source.', ''),
              ('IFCT 2017: Longvah T, Ananthan R, Bhaskarachary K, Venkaiah K. Indian Food Composition Tables 2017. National Institute of Nutrition (ICMR), Hyderabad.', ''),
              ('Data file from the @ifct2017/compositions package (MIT licence, Subhajit Sahu).', '')]
    for i, (t, kind) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name='Arial', size=18, bold=True, color=NAVY) if kind == 'title' else Font(name='Arial', size=10, color=SKY, bold=True) if kind == 'sub' \
            else Font(name='Arial', size=11, bold=True, color=NAVY) if kind == 'h' else base
    ws.column_dimensions['A'].width = 130

    def table(title, data, cols):
        s = wb.create_sheet(title)
        for j, col in enumerate(cols, 1):
            c = s.cell(row=1, column=j, value=col)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical='center', wrap_text=True)
            s.column_dimensions[get_column_letter(j)].width = widths.get(col, 12)
        for i, r in enumerate(data, 2):
            for j, col in enumerate(cols, 1):
                v = r.get(col)
                c = s.cell(row=i, column=j, value=v if v != '' else None)
                c.font = bold if col == 'Food' else base
                c.border = thin
                if isinstance(v, float):
                    c.number_format = '0.0'
            if i % 2 == 0:
                for j in range(1, len(cols) + 1):
                    s.cell(row=i, column=j).fill = PatternFill('solid', fgColor=CREAM)
        s.row_dimensions[1].height = 30
        s.freeze_panes = 'C2'
        s.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(data) + 1}'
        return s

    main = [c for c in COLS if c not in ('Sugar (g)', 'Sat. fat (g)', 'Cholesterol (mg)', 'Calcium (mg)', 'Iron (mg)', 'Sodium (mg)', 'Potassium (mg)', 'Vitamin C (mg)')]
    table('All Foods', rows, main)
    table('Indian Foods (IFCT)', [r for r in rows if r['Source'].startswith('IFCT')],
          [c for c in COLS if c not in ('Household measure', 'Recipe no.', 'Best for')])
    table('Hindivine Dishes', [r for r in rows if r['Source'] == 'Hindivine recipe series'],
          ['ID', 'Recipe no.', 'Food', 'Category', 'Diet', 'Household measure', 'Energy (kcal)', 'Protein (g)', 'Carbs (g)', 'Fat (g)', 'Fibre (g)', 'Best for'])
    table('Ingredients', [r for r in rows if r['Source'] == 'Hindivine ingredient list'],
          ['ID', 'Food', 'Category', 'Diet', 'Household measure', 'Energy (kcal)', 'Protein (g)', 'Carbs (g)', 'Fat (g)', 'Fibre (g)'])

    # Portion calculator: formulas look the food up in All Foods
    pc = wb.create_sheet('Portion Calculator')
    n = len(rows) + 1
    rng = lambda col: f"'All Foods'!${col}$2:${col}${n}"
    letters = {c: get_column_letter(i + 1) for i, c in enumerate(main)}
    pc['A1'] = 'Portion Calculator'
    pc['A1'].font = Font(name='Arial', size=16, bold=True, color=NAVY)
    pc['A2'] = 'Type a food ID from the All Foods sheet in column A (yellow) and the amount in column B (grams for "per 100 g" foods, number of servings for dishes).'
    pc['A2'].font = Font(name='Arial', size=10, italic=True, color='7F7068')
    hdr = ['Food ID', 'Amount', 'Food', 'Basis', 'Energy (kcal)', 'Protein (g)', 'Carbs (g)', 'Fat (g)', 'Fibre (g)']
    for j, h in enumerate(hdr, 1):
        c = pc.cell(row=4, column=j, value=h)
        c.font, c.fill = head_font, head_fill
        pc.column_dimensions[get_column_letter(j)].width = [12, 10, 44, 13, 13, 12, 12, 10, 10][j - 1]
    examples = [('IF-A003', 100), (next(r['ID'] for r in rows if r['Food'] == 'Paneer' and r['Source'] == 'Hindivine ingredient list'), 50), ('HR-0002', 1), ('IF-E028', 150)]
    yellow = PatternFill('solid', fgColor='FFF59D')
    for i in range(5, 25):
        a = pc.cell(row=i, column=1)
        b = pc.cell(row=i, column=2)
        if i - 5 < len(examples):
            a.value, b.value = examples[i - 5]
        for c in (a, b):
            c.fill, c.font = yellow, Font(name='Arial', size=10, color='0000FF')
        m = f'MATCH($A{i},{rng("A")},0)'
        factor = f'IF(D{i}="per 100 g",$B{i}/100,$B{i})'
        pc.cell(row=i, column=3, value=f'=IF($A{i}="","",IFERROR(INDEX({rng(letters["Food"])},{m}),"ID not found"))')
        pc.cell(row=i, column=4, value=f'=IF($A{i}="","",IFERROR(INDEX({rng(letters["Basis"])},{m}),""))')
        for j, col in enumerate(['Energy (kcal)', 'Protein (g)', 'Carbs (g)', 'Fat (g)', 'Fibre (g)'], 5):
            c = pc.cell(row=i, column=j, value=f'=IF(OR($A{i}="",$D{i}=""),"",IFERROR(INDEX({rng(letters[col])},{m})*{factor},""))')
            c.number_format = '0' if j == 5 else '0.0'
        for j in range(3, 10):
            pc.cell(row=i, column=j).font = base
    pc.cell(row=25, column=3, value='Total').font = bold
    for j in range(5, 10):
        L = get_column_letter(j)
        c = pc.cell(row=25, column=j, value=f'=SUM({L}5:{L}24)')
        c.font, c.number_format = bold, '0' if j == 5 else '0.0'
    pc['A27'] = 'Example rows: bajra (100 g), paneer (50 g), Avocado Chila (1 serving), white guava (150 g). Replace them with your own.'
    pc['A27'].font = Font(name='Arial', size=9, italic=True, color='7F7068')
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    pc.add_data_validation(dv)
    dv.add('B5:B24')
    pc.freeze_panes = 'A5'
    wb.move_sheet('Portion Calculator', offset=-(len(wb.sheetnames) - 2))
    wb.save(path)


def write_html(rows, path):
    tpl = open(os.path.join(HERE, 'foodfinder_template.html'), encoding='utf-8').read()
    keys = ['ID', 'Food', 'Hindi / local name', 'Category', 'Source', 'Diet', 'Basis', 'Household measure', 'Energy (kcal)', 'Protein (g)',
            'Carbs (g)', 'Fat (g)', 'Fibre (g)', 'Sugar (g)', 'Calcium (mg)', 'Iron (mg)', 'Sodium (mg)', 'Vitamin C (mg)', 'Recipe no.', 'Best for']
    data = {'cols': keys, 'rows': [[r.get(k) if r.get(k) != '' else None for k in keys] + [r.get('_local', '')] for r in rows]}
    import base64
    logo = base64.b64encode(open(os.path.join(HERE, 'logo-small.jpg'), 'rb').read()).decode()
    html = tpl.replace('/*__DATA__*/null', json.dumps(data, ensure_ascii=False, separators=(',', ':'))).replace('__LOGO__', 'data:image/jpeg;base64,' + logo)
    if '--artifact' not in sys.argv:
        t = html.index('</title>') + len('</title>')
        e = html.index('</style>') + 8
        html = ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
                '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
                + html[:t] + html[t:e] + '\n</head>\n<body>\n' + html[e:] + '\n</body>\n</html>\n')
    with open(path, 'w', encoding='utf-8') as fh:
        fh.write(html)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else '.'
    os.makedirs(out, exist_ok=True)
    rows = all_rows()
    write_csv(rows, os.path.join(out, 'Hindivine-Food-List.csv'))
    write_xlsx(rows, os.path.join(out, 'Hindivine-Food-List.xlsx'))
    if os.path.exists(os.path.join(HERE, 'foodfinder_template.html')):
        write_html(rows, os.path.join(out, 'Hindivine-Food-Finder.html'))
    by = {}
    for r in rows:
        by[r['Source']] = by.get(r['Source'], 0) + 1
    print(len(rows), 'foods', by)


if __name__ == '__main__':
    main()
